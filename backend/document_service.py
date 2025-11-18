from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from typing import List
import models
from sqlalchemy.orm import Session


class DocumentExporter:
    
    @staticmethod
    def export_docx(project: models.Project, sections: List[models.DocumentSection]) -> bytes:
        """Export project content as a .docx file."""
        
        doc = Document()
        
        # Add title
        title = doc.add_heading(project.title, 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Add main topic as subtitle
        subtitle = doc.add_paragraph(project.main_topic)
        subtitle_format = subtitle.paragraph_format
        subtitle_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        subtitle.runs[0].font.size = Pt(14)
        subtitle.runs[0].font.color.rgb = RGBColor(128, 128, 128)
        
        doc.add_paragraph()  # Add spacing
        
        # Add sections in order
        sorted_sections = sorted(sections, key=lambda x: x.order)
        
        for section in sorted_sections:
            # Add section heading
            doc.add_heading(section.title, 1)
            
            # Add section content
            if section.content:
                # Split content into paragraphs
                paragraphs = section.content.split('\n\n')
                for para in paragraphs:
                    if para.strip():
                        p = doc.add_paragraph(para.strip())
                        p.paragraph_format.line_spacing = 1.5
            else:
                doc.add_paragraph("[Content not yet generated]")
            
            doc.add_paragraph()  # Add spacing between sections
        
        # Save to bytes
        from io import BytesIO
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    
    @staticmethod
    def export_pptx(project: models.Project, sections: List[models.DocumentSection]) -> bytes:
        """Export project content as a .pptx file."""
        
        prs = Presentation()
        prs.slide_width = PptxInches(10)
        prs.slide_height = PptxInches(7.5)
        
        # Title slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = project.title
        subtitle.text = project.main_topic
        
        # Add content slides
        sorted_sections = sorted(sections, key=lambda x: x.order)
        
        for section in sorted_sections:
            # Use bullet slide layout
            bullet_slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(bullet_slide_layout)
            
            # Set title
            title = slide.shapes.title
            title.text = section.title
            
            # Set content
            body_shape = slide.shapes.placeholders[1]
            text_frame = body_shape.text_frame
            text_frame.clear()
            
            if section.content:
                # Split content by lines
                lines = section.content.strip().split('\n')
                for i, line in enumerate(lines):
                    line = line.strip()
                    if line:
                        # Remove leading dashes or bullets
                        line = line.lstrip('-•* ')
                        
                        if i == 0:
                            p = text_frame.paragraphs[0]
                            p.text = line
                        else:
                            p = text_frame.add_paragraph()
                            p.text = line
                        
                        p.level = 0
                        p.font.size = PptxPt(18)
            else:
                p = text_frame.paragraphs[0]
                p.text = "[Content not yet generated]"
        
        # Save to bytes
        from io import BytesIO
        buffer = BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


document_exporter = DocumentExporter()
